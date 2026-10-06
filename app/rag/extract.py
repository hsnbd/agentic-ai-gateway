"""Turn an uploaded file into the plain text that gets chunked and embedded.

Text, Markdown, and HTML need nothing extra. PDF and Word (.docx) use the
optional ``rag-docs`` extra (``pypdf``, ``python-docx``), which the Docker
image installs. Unsupported or unreadable files are refused with a clear
error rather than ingested as garbage.
"""

from __future__ import annotations

import io
from html.parser import HTMLParser
from pathlib import PurePath

from app.core.errors import InvalidRequestError

#: File extension -> stored content type.
MEDIA_TYPES = {
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".markdown": "text/markdown",
    ".html": "text/html",
    ".htm": "text/html",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}

_SKIPPED_TAGS = {"script", "style", "noscript", "template", "head"}
_BLOCK_TAGS = {
    "p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6",
    "section", "article", "header", "footer", "blockquote", "pre", "table",
}  # fmt: skip


def extract_text(data: bytes, filename: str) -> tuple[str, str]:
    """Return ``(text, content_type)`` for an uploaded file."""
    suffix = PurePath(filename).suffix.lower()
    media_type = MEDIA_TYPES.get(suffix)
    if media_type is None:
        raise InvalidRequestError(
            f"Cannot ingest {suffix or 'files without an extension'}; supported: "
            + ", ".join(sorted(MEDIA_TYPES)),
            status_code=415,
        )
    if suffix == ".pdf":
        text = _pdf_text(data)
    elif suffix == ".docx":
        text = _docx_text(data)
    elif media_type == "text/html":
        text = html_to_text(data.decode("utf-8", errors="replace"))
    else:
        text = data.decode("utf-8", errors="replace")
    if not text.strip():
        raise InvalidRequestError(
            f"No text could be extracted from {filename!r} (a scanned PDF needs OCR first)",
            status_code=422,
        )
    return text, media_type


class _HtmlText(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._skipping = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in _SKIPPED_TAGS:
            self._skipping += 1
        elif tag in _BLOCK_TAGS:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in _SKIPPED_TAGS:
            self._skipping = max(0, self._skipping - 1)
        elif tag in _BLOCK_TAGS:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skipping:
            self.parts.append(data)


def html_to_text(html: str) -> str:
    """Visible text of an HTML page, one block per line; scripts and styles dropped."""
    parser = _HtmlText()
    parser.feed(html)
    parser.close()
    lines = (" ".join(line.split()) for line in "".join(parser.parts).splitlines())
    return "\n".join(line for line in lines if line)


def _missing_extra(kind: str) -> InvalidRequestError:  # pragma: no cover - extra installed
    return InvalidRequestError(
        f"{kind} ingestion needs the rag-docs extra (pip install 'aigateway[rag-docs]')",
        status_code=415,
    )


def _pdf_text(data: bytes) -> str:
    try:
        import pypdf
    except ImportError:  # pragma: no cover - the dev and Docker installs include it
        raise _missing_extra("PDF") from None
    try:
        reader = pypdf.PdfReader(io.BytesIO(data))
        return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages)
    except Exception as exc:
        raise InvalidRequestError(f"Could not read the PDF: {exc}", status_code=422) from exc


def _docx_text(data: bytes) -> str:
    try:
        import docx
    except ImportError:  # pragma: no cover - the dev and Docker installs include it
        raise _missing_extra("Word") from None
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise InvalidRequestError(f"Could not read the Word file: {exc}", status_code=422) from exc
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(part for part in parts if part.strip())
